import sqlite3
import os
import uuid
import qrcode
from datetime import datetime, timedelta
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DB_PATH = os.path.join(os.path.dirname(__file__), "smartpass.db")
STATIC_DIR = os.path.join(os.path.dirname(__file__), "static")
QR_DIR = os.path.join(STATIC_DIR, "qrcodes")
RECEIPTS_DIR = os.path.join(STATIC_DIR, "receipts")

os.makedirs(QR_DIR, exist_ok=True)
os.makedirs(RECEIPTS_DIR, exist_ok=True)

def get_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_connection()
    c = conn.cursor()
    
    # Users
    c.execute("""
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        telegram_id INTEGER UNIQUE,
        username TEXT,
        full_name TEXT NOT NULL,
        phone TEXT,
        client_type TEXT DEFAULT 'individual',
        created_at TEXT
    )
    """)
    
    # Products / Services
    c.execute("""
    CREATE TABLE IF NOT EXISTS products (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        network_name TEXT NOT NULL,
        service_name TEXT NOT NULL,
        denomination TEXT NOT NULL,
        price_legal REAL NOT NULL,
        price_fop REAL NOT NULL,
        price_usdt REAL NOT NULL,
        is_active INTEGER DEFAULT 1
    )
    """)
    
    # Orders
    c.execute("""
    CREATE TABLE IF NOT EXISTS orders (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        order_number TEXT UNIQUE NOT NULL,
        user_id INTEGER NOT NULL,
        product_id INTEGER NOT NULL,
        quantity INTEGER NOT NULL,
        total_amount REAL NOT NULL,
        currency TEXT NOT NULL,
        payment_method TEXT NOT NULL,
        receipt_file TEXT,
        status TEXT DEFAULT 'NEW',
        admin_notes TEXT,
        created_at TEXT,
        approved_at TEXT,
        FOREIGN KEY (user_id) REFERENCES users(id),
        FOREIGN KEY (product_id) REFERENCES products(id)
    )
    """)
    
    # Vouchers / Passes
    c.execute("""
    CREATE TABLE IF NOT EXISTS vouchers (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        voucher_code TEXT UNIQUE NOT NULL,
        order_id INTEGER NOT NULL,
        user_id INTEGER NOT NULL,
        network_name TEXT NOT NULL,
        service_name TEXT NOT NULL,
        denomination TEXT NOT NULL,
        status TEXT DEFAULT 'AVAILABLE',
        qr_image TEXT,
        created_at TEXT,
        expires_at TEXT,
        used_at TEXT,
        FOREIGN KEY (order_id) REFERENCES orders(id),
        FOREIGN KEY (user_id) REFERENCES users(id)
    )
    """)
    
    # Check-in QR requests
    c.execute("""
    CREATE TABLE IF NOT EXISTS checkins (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        voucher_id INTEGER NOT NULL,
        user_id INTEGER NOT NULL,
        network_name TEXT NOT NULL,
        qr_token TEXT UNIQUE NOT NULL,
        status TEXT DEFAULT 'PENDING',
        created_at TEXT,
        expires_at TEXT,
        FOREIGN KEY (voucher_id) REFERENCES vouchers(id),
        FOREIGN KEY (user_id) REFERENCES users(id)
    )
    """)
    
    # Audit trail
    c.execute("""
    CREATE TABLE IF NOT EXISTS audit_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        admin_name TEXT NOT NULL,
        action TEXT NOT NULL,
        details TEXT NOT NULL,
        ip_address TEXT,
        created_at TEXT
    )
    """)
    
    conn.commit()
    conn.close()

def log_audit(admin_name, action, details, ip_address="127.0.0.1"):
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
    INSERT INTO audit_logs (admin_name, action, details, ip_address, created_at)
    VALUES (?, ?, ?, ?, ?)
    """, (admin_name, action, details, ip_address, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
    conn.commit()
    conn.close()

def generate_qr(token_data, filename):
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=3,
    )
    qr.add_data(token_data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#1E293B", back_color="white")
    filepath = os.path.join(QR_DIR, filename)
    img.save(filepath)
    return f"/static/qrcodes/{filename}"

def approve_order(order_id, admin_name="Головний Адміністратор"):
    conn = get_connection()
    c = conn.cursor()
    
    # Check current order
    c.execute("""
    SELECT o.*, p.network_name, p.service_name, p.denomination, u.full_name
    FROM orders o
    JOIN products p ON o.product_id = p.id
    JOIN users u ON o.user_id = u.id
    WHERE o.id = ?
    """, (order_id,))
    order = c.fetchone()
    
    if not order or order["status"] == "APPROVED":
        conn.close()
        return False
        
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    expires_str = (datetime.now() + timedelta(days=90)).strftime("%Y-%m-%d")
    
    # Generate vouchers atomically
    vouchers_created = []
    for i in range(order["quantity"]):
        code = f"SP-{order['network_name'][:3].upper()}-{uuid.uuid4().hex[:6].upper()}"
        qr_file = f"{code}.png"
        qr_rel_path = generate_qr(f"SMARTPASS_VERIFY:{code}", qr_file)
        
        c.execute("""
        INSERT INTO vouchers (voucher_code, order_id, user_id, network_name, service_name, denomination, status, qr_image, created_at, expires_at)
        VALUES (?, ?, ?, ?, ?, ?, 'AVAILABLE', ?, ?, ?)
        """, (code, order["id"], order["user_id"], order["network_name"], order["service_name"], order["denomination"], qr_rel_path, now_str, expires_str))
        vouchers_created.append(code)
        
    c.execute("""
    UPDATE orders 
    SET status = 'APPROVED', approved_at = ?, admin_notes = ?
    WHERE id = ?
    """, (now_str, f"Схвалено адміністратором {admin_name}. Випущено талонів: {len(vouchers_created)}", order_id))
    
    conn.commit()
    conn.close()
    
    log_audit(
        admin_name, 
        "СХВАЛЕННЯ_ОПЛАТИ", 
        f"Замовлення #{order['order_number']} схвалено. Клієнт: {order['full_name']}. Видано талони: {', '.join(vouchers_created)}"
    )
    return True

def reject_order(order_id, reason="Некоректна сума або нечитабельний чек", admin_name="Головний Адміністратор"):
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT * FROM orders WHERE id = ?", (order_id,))
    order = c.fetchone()
    if not order:
        conn.close()
        return False
        
    c.execute("""
    UPDATE orders 
    SET status = 'REJECTED', admin_notes = ? 
    WHERE id = ?
    """, (reason, order_id))
    conn.commit()
    conn.close()
    
    log_audit(admin_name, "ВІДХИЛЕННЯ_ОПЛАТИ", f"Замовлення #{order['order_number']} відхилено. Причина: {reason}")
    return True

def confirm_checkin(checkin_id, admin_name="Адміністратор рецепції"):
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
    SELECT ch.*, v.voucher_code, v.denomination, u.full_name
    FROM checkins ch
    JOIN vouchers v ON ch.voucher_id = v.id
    JOIN users u ON ch.user_id = u.id
    WHERE ch.id = ?
    """, (checkin_id,))
    ch = c.fetchone()
    
    if not ch or ch["status"] != "PENDING":
        conn.close()
        return False
        
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    # Atomic transaction
    c.execute("UPDATE checkins SET status = 'CONFIRMED' WHERE id = ?", (checkin_id,))
    c.execute("UPDATE vouchers SET status = 'USED', used_at = ? WHERE id = ?", (now_str, ch["voucher_id"]))
    
    conn.commit()
    conn.close()
    
    log_audit(
        admin_name,
        "ПОГАШЕННЯ_ТАЛОНА",
        f"Талон {ch['voucher_code']} ({ch['network_name']} - {ch['denomination']}) успішно погашено для клієнта {ch['full_name']}."
    )
    return True

def cancel_checkin(checkin_id, admin_name="Адміністратор рецепції"):
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT * FROM checkins WHERE id = ?", (checkin_id,))
    ch = c.fetchone()
    if not ch or ch["status"] != "PENDING":
        conn.close()
        return False
        
    c.execute("UPDATE checkins SET status = 'CANCELLED' WHERE id = ?", (checkin_id,))
    c.execute("UPDATE vouchers SET status = 'AVAILABLE' WHERE id = ?", (ch["voucher_id"],))
    conn.commit()
    conn.close()
    
    log_audit(admin_name, "СКАСУВАННЯ_ЧЕКІНУ", f"Запит на чек-ін #{checkin_id} скасовано. Талон повернено на баланс.")
    return True

def seed_demo_data():
    init_db()
    conn = get_connection()
    c = conn.cursor()
    
    # Check if products already exist
    c.execute("SELECT COUNT(*) FROM products")
    if c.fetchone()[0] > 0:
        conn.close()
        return
        
    # Products
    products = [
        ("SpaceHub Coworking", "Денний абонемент (Day Pass)", "1 день (Open Space + Lounge)", 650.0, 550.0, 13.5),
        ("SpaceHub Coworking", "Пакет візитів Flex 10", "10 візитів (дійсний 60 днів)", 5400.0, 4800.0, 118.0),
        ("Grand SPA Resort", "СПА-день All-Inclusive", "Басейн, сауни, джакузі, рушник", 1600.0, 1400.0, 34.0),
        ("Grand SPA Resort", "Подарунковий сертифікат Relax", "Номінал 2 500 грн на всі послуги", 2700.0, 2500.0, 61.0),
        ("Titan Fitness Club", "Абонемент місячний Full Day", "Тренажерний зал + групові заняття", 2200.0, 1950.0, 48.0),
        ("Titan Fitness Club", "Клубна карта 10 тренувань", "10 тренувань з доступом до СПА", 3100.0, 2800.0, 68.0),
    ]
    for p in products:
        c.execute("""
        INSERT INTO products (network_name, service_name, denomination, price_legal, price_fop, price_usdt)
        VALUES (?, ?, ?, ?, ?, ?)
        """, p)
        
    # Users
    users = [
        (10948291, "oleg_tech", "Олег Мельник", "+380501234567", "fop"),
        (28391029, "kateryna_m", "Катерина Білоус", "+380679876543", "individual"),
        (94820194, "dmytro_dev", "Дмитро Коваленко (ТОВ 'ДатаЛінк')", "+380631122334", "legal"),
    ]
    for u in users:
        c.execute("""
        INSERT INTO users (telegram_id, username, full_name, phone, client_type, created_at)
        VALUES (?, ?, ?, ?, ?, ?)
        """, (*u, (datetime.now() - timedelta(days=5)).strftime("%Y-%m-%d %H:%M:%S")))
        
    conn.commit()
    
    # Create sample orders
    sample_orders = [
        ("ORD-2026-1001", 1, 2, 1, 4800.0, "UAH", "FOP", "PAID_REPORTED", "receipt_sample_1.png", (datetime.now() - timedelta(hours=2)).strftime("%Y-%m-%d %H:%M:%S")),
        ("ORD-2026-1002", 2, 4, 1, 61.0, "USDT", "USDT", "PAID_REPORTED", "receipt_sample_2.png", (datetime.now() - timedelta(minutes=45)).strftime("%Y-%m-%d %H:%M:%S")),
        ("ORD-2026-0988", 3, 1, 5, 3250.0, "UAH", "LEGAL", "APPROVED", "receipt_sample_3.png", (datetime.now() - timedelta(days=2)).strftime("%Y-%m-%d %H:%M:%S")),
    ]
    
    for so in sample_orders:
        c.execute("""
        INSERT INTO orders (order_number, user_id, product_id, quantity, total_amount, currency, payment_method, status, receipt_file, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, so)
        
    conn.commit()
    conn.close()
    
    # Approve order 3 to generate initial vouchers
    approve_order(3, "Система (Demo Seed)")
    
    # Create a live check-in request for demo
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT id, user_id, network_name FROM vouchers WHERE status = 'AVAILABLE' LIMIT 1")
    v = c.fetchone()
    if v:
        c.execute("UPDATE vouchers SET status = 'RESERVED' WHERE id = ?", (v["id"],))
        token = f"CHK-{uuid.uuid4().hex[:8].upper()}"
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        exp_str = (datetime.now() + timedelta(minutes=15)).strftime("%Y-%m-%d %H:%M:%S")
        c.execute("""
        INSERT INTO checkins (voucher_id, user_id, network_name, qr_token, status, created_at, expires_at)
        VALUES (?, ?, ?, ?, 'PENDING', ?, ?)
        """, (v["id"], v["user_id"], v["network_name"], token, now_str, exp_str))
        conn.commit()
    conn.close()

def generate_excel_report():
    wb = openpyxl.Workbook()
    
    # Header styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    conn = get_connection()
    c = conn.cursor()
    
    # Sheet 1: Замовлення
    ws_orders = wb.active
    ws_orders.title = "Замовлення"
    headers_orders = ["№ Замовлення", "Клієнт", "Телефон", "Послуга / Мережа", "К-сть", "Сума", "Валюта", "Оплата", "Статус", "Дата створення"]
    ws_orders.append(headers_orders)
    
    c.execute("""
    SELECT o.order_number, u.full_name, u.phone, p.service_name || ' (' || p.network_name || ')', 
           o.quantity, o.total_amount, o.currency, o.payment_method, o.status, o.created_at
    FROM orders o
    JOIN users u ON o.user_id = u.id
    JOIN products p ON o.product_id = p.id
    ORDER BY o.id DESC
    """)
    for row in c.fetchall():
        ws_orders.append(list(row))
        
    # Sheet 2: Талони та Абонементи
    ws_vouchers = wb.create_sheet(title="Талони")
    headers_v = ["Код талона", "Мережа", "Послуга", "Номінал", "Клієнт", "Статус", "Дата випуску", "Дійсний до", "Дата використання"]
    ws_vouchers.append(headers_v)
    
    c.execute("""
    SELECT v.voucher_code, v.network_name, v.service_name, v.denomination, u.full_name, v.status, v.created_at, v.expires_at, v.used_at
    FROM vouchers v
    JOIN users u ON v.user_id = u.id
    ORDER BY v.id DESC
    """)
    for row in c.fetchall():
        ws_vouchers.append(list(row))
        
    # Sheet 3: Журнал аудиту
    ws_audit = wb.create_sheet(title="Журнал дій")
    headers_a = ["ID", "Адміністратор", "Дія", "Деталі", "IP-адреса", "Час події"]
    ws_audit.append(headers_a)
    
    c.execute("SELECT id, admin_name, action, details, ip_address, created_at FROM audit_logs ORDER BY id DESC")
    for row in c.fetchall():
        ws_audit.append(list(row))
        
    conn.close()
    
    # Apply styling to all sheets
    for ws in [ws_orders, ws_vouchers, ws_audit]:
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = thin_border
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
